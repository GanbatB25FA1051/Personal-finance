"""Khan Bank Excel statement parser."""
import re
from datetime import datetime, date
from decimal import InvalidOperation
from openpyxl.workbook import Workbook
from banks.base import BankParser, REGISTRY
from core.models import ParsedRow, ParsedStatement, to_minor

HEADER_SEARCH_ROWS = 10
HEADERS = ['Гүйлгээний огноо', 'Салбар', 'Эхний үлдэгдэл', 'Кредит гүйлгээ',
           'Дебит гүйлгээ', 'Эцсийн үлдэгдэл', 'Гүйлгээний утга', 'Харьцсан данс']


def text(value: object) -> str:
    """Normalize optional cell text."""
    return '' if value is None else str(value).strip()


def iso_date(value: object, full: bool = False) -> str:
    """Convert Excel dates and slash-separated date strings."""
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        parsed = datetime.fromisoformat(text(value).replace('/', '-'))
    return parsed.strftime('%Y-%m-%d %H:%M:%S' if full else '%Y-%m-%d')


class KhanBankParser(BankParser):
    bank_code = 'khan'

    def can_parse(self, wb: Workbook) -> bool:
        """Find identifying headers within the supported header-search window."""
        return any({'Харьцсан данс', 'Гүйлгээний огноо'} <= {text(v) for v in row}
                   for row in wb.active.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS, values_only=True))

    def parse(self, wb: Workbook) -> ParsedStatement:
        """Parse by column names; report invalid rows without aborting."""
        sheet = wb.active
        metadata = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        meta = lambda i: metadata[i] if i < len(metadata) else None
        account = re.sub(r'\s+', '', re.sub(r'^IBAN\s*:', '', text(meta(11)), flags=re.I))
        currency = next((text(v).upper() for v in metadata if re.fullmatch(r'[A-Za-z]{3}', text(v))), '')
        result = ParsedStatement(
            self.bank_code, account, currency, None, None, [], None, None, [],
            rows_skipped=0, fatal_errors=[],
        )
        if not account:
            result.fatal_errors.append('Account number was not found')
        if not currency:
            result.fatal_errors.append('Currency code was not found')
        for attr, index in [('period_start', 6), ('period_end', 8)]:
            try:
                if meta(index):
                    setattr(result, attr, iso_date(meta(index)))
            except (ValueError, TypeError):
                result.warnings.append(f'Invalid {attr} metadata')
        columns = None
        credit_sum = debit_sum = 0
        for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = [text(v) for v in row]
            if columns is None:
                if number <= HEADER_SEARCH_ROWS and {'Харьцсан данс', 'Гүйлгээний огноо'} <= set(values):
                    if not set(HEADERS) <= set(values):
                        result.fatal_errors.append('Missing required statement columns')
                        return result
                    columns = {name: values.index(name) for name in HEADERS}
                continue
            if not any(values):
                continue
            cell = lambda name: row[columns[name]] if columns[name] < len(row) else None
            try:
                if any(value.startswith('Нийт дүн:') for value in values):
                    credit_value = cell('Кредит гүйлгээ')
                    debit_value = cell('Дебит гүйлгээ')
                    try:
                        result.control_credit = to_minor(credit_value if text(credit_value) else '0')
                        result.control_debit = -abs(to_minor(debit_value if text(debit_value) else '0'))
                    except (ValueError, TypeError, InvalidOperation, OverflowError):
                        result.warnings.append('Statement control totals could not be read')
                    continue

                money = lambda name: to_minor(cell(name) if text(cell(name)) else '0')
                timestamp = iso_date(cell('Гүйлгээний огноо'), True)
                credit = money('Кредит гүйлгээ')
                raw_debit = money('Дебит гүйлгээ')
                debit = -abs(raw_debit)
                if raw_debit > 0:
                    result.warnings.append(
                        f'Row {number}: positive debit normalized to a negative amount'
                    )
                opening, closing = money('Эхний үлдэгдэл'), money('Эцсийн үлдэгдэл')
                amount = credit + debit
                result.rows.append(ParsedRow(timestamp, amount, closing, text(cell('Гүйлгээний утга')),
                                             text(cell('Харьцсан данс')) or None, text(cell('Салбар')) or None))
                credit_sum += credit
                debit_sum += debit
                if abs(opening + amount - closing) > 1:
                    result.warnings.append(f'Row {number}: opening balance plus amount differs from closing balance')
            except (ValueError, TypeError, InvalidOperation, OverflowError):
                result.rows_skipped += 1
                result.warnings.append(f'Row {number}: invalid statement data; row skipped')
        if columns is None:
            result.fatal_errors.append('Statement header not found')
        for label, actual, control in [('credit', credit_sum, result.control_credit),
                                        ('debit', debit_sum, result.control_debit)]:
            if control is not None and abs(actual - control) > 1:
                result.warnings.append(f'Total {label} differs from statement control total')
        return result

    def is_fee(self, raw_description: str) -> bool:
        """Recognize supported fee descriptions."""
        raw = raw_description.casefold().strip()
        return 'хураамж' in raw or 'fino fee' in raw or (raw.startswith('qpay') and raw.endswith('шимтгэл'))

    def clean_merchant(self, raw_description: str) -> str:
        """Strip QPay identifiers while retaining merchant names."""
        raw = raw_description.casefold().strip()
        if re.match(r'^z,\s*\d+', raw):
            return 'z-recurring'
        if raw.startswith('qpay'):
            if ',' in raw:
                parts = [p.strip() for p in raw.split(',')]
                raw = next((p for p in reversed(parts) if p and p not in {'qpay', 'qmenu', 'sms', 'charge'}
                            and not re.fullmatch(r'[\d\s-]+', p)), '')
            raw = re.sub(r'^qpay\s+(?:fino(?: fee)?-)?\d+\s*', '', raw)
            raw = re.sub(r'\b(?:trx|sms|charge)\s*:.*$', '', raw)
            raw = re.sub(r'\s+\d+(?:[.,]\d+)?\s*$', '', raw)
            raw = re.sub(r'\b(?:qpay|qmenu|sms|charge)\b', '', raw)
        raw = re.sub(r'\b\d{8,}\b', '', raw)
        return re.sub(r'\s+', ' ', raw).strip(' ,;:.-')


REGISTRY.append(KhanBankParser())

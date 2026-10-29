"""Exercise complete imports, overlapping files, and atomic rollback."""
import pytest
from openpyxl import load_workbook
from core.db import get_conn, init_db
from core.ingest import (
    import_preview,
    import_statement,
    preview_statement,
    statement_account_number,
)
from core.classify import set_category
from core.reports import monthly_summary
from tests.make_fixture import make_fixture


def test_import_and_duplicates(tmp_path) -> None:
    path = tmp_path / 'test.xlsx'
    make_fixture(path)
    assert statement_account_number(path.read_bytes()) == 'MN030005005000000001'
    conn = get_conn(':memory:')
    init_db(conn)
    first = import_statement(conn, path, path.name)
    assert first.rows_new == 12
    archived = conn.execute(
        'SELECT file_data, file_hash FROM import_files WHERE import_id=?',
        (first.import_id,),
    ).fetchone()
    assert archived['file_data'] == path.read_bytes()
    assert archived['file_hash']
    first_links = conn.execute(
        'SELECT COUNT(*) FROM import_transactions WHERE import_id=?',
        (first.import_id,),
    ).fetchone()[0]
    assert first_links == 12
    fees = conn.execute(
        """SELECT * FROM transactions
           WHERE category_id=(SELECT id FROM categories
                              WHERE name='Bank Fees & Charges')"""
    ).fetchall()
    assert len(fees) == 3
    assert all(r['parent_txn_id'] is not None for r in fees)
    assert first.auto_classified == 9
    assert first.needs_review == 0
    second = import_statement(conn, path.read_bytes(), path.name)
    assert (second.rows_new, second.rows_duplicate) == (0, 12)
    assert 'This exact file was already imported' in second.warnings
    second_links = conn.execute(
        'SELECT COUNT(*) FROM import_transactions WHERE import_id=?',
        (second.import_id,),
    ).fetchone()[0]
    assert second_links == 12
    wb = load_workbook(path)
    ws = wb.active
    ws.delete_rows(13, 2)
    for i in range(9, 13):
        ws.cell(i, 1, f'2026-11-{i:02d} 12:00:00')
    wb.save(path)
    wb.close()
    overlap = import_statement(conn, path, path.name)
    assert (overlap.rows_new, overlap.rows_duplicate) == (4, 6)
    overlap_links = conn.execute(
        'SELECT COUNT(*) FROM import_transactions WHERE import_id=?',
        (overlap.import_id,),
    ).fetchone()[0]
    assert overlap_links == 10
    conn.close()


def test_next_period_statement_appends_to_same_account_history(tmp_path) -> None:
    first_path = tmp_path / 'october.xlsx'
    next_path = tmp_path / 'november.xlsx'
    make_fixture(first_path)
    make_fixture(next_path)

    workbook = load_workbook(next_path)
    sheet = workbook.active
    for row_number in range(3, sheet.max_row + 1):
        value = sheet.cell(row_number, 1).value
        if isinstance(value, str) and value.startswith('2026-10-'):
            sheet.cell(row_number, 1, value.replace('2026-10-', '2026-11-', 1))
    workbook.save(next_path)
    workbook.close()

    conn = get_conn(':memory:')
    init_db(conn)
    try:
        october = import_statement(conn, first_path, first_path.name)
        november = import_statement(conn, next_path, next_path.name)

        assert october.rows_new == november.rows_new == 12
        assert conn.execute('SELECT COUNT(*) FROM accounts').fetchone()[0] == 1
        assert conn.execute('SELECT COUNT(*) FROM transactions').fetchone()[0] == 24
        assert conn.execute('SELECT COUNT(*) FROM imports').fetchone()[0] == 2

        monthly = monthly_summary(conn, '2026-10-01', '2026-11-30')
        assert list(monthly['month']) == ['2026-10', '2026-11']
        assert monthly.loc[0, 'income'] == monthly.loc[1, 'income']
        assert monthly.loc[0, 'expense'] == monthly.loc[1, 'expense']
    finally:
        conn.close()


def test_new_weekly_statement_does_not_learn_rule_from_manual_choice(
    tmp_path,
) -> None:
    path = tmp_path / 'weekly.xlsx'
    make_fixture(path)
    conn = get_conn(':memory:')
    init_db(conn)
    first = import_statement(conn, path, path.name)
    merchant = conn.execute(
        """SELECT id, merchant_clean FROM transactions
           WHERE raw_description LIKE 'qpay fino-940482663759306%'"""
    ).fetchone()
    assert merchant is not None
    category_id = conn.execute(
        "SELECT id FROM categories WHERE name='Groceries / Food'"
    ).fetchone()[0]
    set_category(conn, [merchant['id']], category_id)

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.cell(3, 1, '2026-10-20 12:00:00')
    sheet.cell(3, 7, 'qpay fino-123456789012345, Дэвжих супермаркет')
    workbook.save(path)
    workbook.close()

    second = import_statement(conn, path, 'weekly.xlsx')
    assert second.rows_new == 1
    linked = conn.execute(
        '''SELECT category_id, classified_by FROM transactions
           WHERE txn_datetime='2026-10-20 12:00:00' '''
    ).fetchone()
    other_expense = conn.execute(
        "SELECT id FROM categories WHERE name='Other expense'"
    ).fetchone()[0]
    assert tuple(linked) == (other_expense, 'system')
    assert second.auto_classified == 1
    assert conn.execute(
        '''SELECT COUNT(*) FROM merchant_rules
           WHERE match_type='merchant_exact'
             AND pattern=? COLLATE NOCASE''',
        (merchant['merchant_clean'],),
    ).fetchone()[0] == 0
    conn.close()


def test_invalid_file() -> None:
    conn = get_conn(':memory:')
    init_db(conn)
    with pytest.raises(ValueError, match='Unrecognized'):
        import_statement(conn, b'broken', 'bad.xlsx')
    conn.close()


def test_rollback(tmp_path, monkeypatch) -> None:
    from core import classify
    path = tmp_path / 'test.xlsx'
    make_fixture(path)
    conn = get_conn(':memory:')
    init_db(conn)
    def fail(*args):
        raise RuntimeError('forced failure')
    monkeypatch.setattr(classify, 'apply_rules', fail)
    with pytest.raises(RuntimeError):
        import_statement(conn, path, path.name)
    for table in ('accounts', 'imports', 'transactions'):
        assert conn.execute('SELECT COUNT(*) FROM ' + table).fetchone()[0] == 0
    conn.close()


def test_commit_failure_rolls_back_import_and_preserves_original_error(
    tmp_path,
) -> None:
    path = tmp_path / 'commit-failure.xlsx'
    make_fixture(path)
    inner = get_conn(':memory:')
    init_db(inner)

    class CommitFails:
        @property
        def in_transaction(self):
            return inner.in_transaction

        def execute(self, *args, **kwargs):
            return inner.execute(*args, **kwargs)

        def commit(self):
            raise RuntimeError('commit failed')

        def rollback(self):
            return inner.rollback()

    try:
        with pytest.raises(RuntimeError, match='commit failed'):
            import_statement(CommitFails(), path, path.name)
        for table in ('accounts', 'imports', 'transactions'):
            assert inner.execute(
                'SELECT COUNT(*) FROM ' + table
            ).fetchone()[0] == 0
    finally:
        inner.close()


def test_import_failure_preserves_callers_prior_transaction(tmp_path, monkeypatch) -> None:
    from core import classify

    path = tmp_path / 'nested-rollback.xlsx'
    make_fixture(path)
    conn = get_conn(':memory:')
    init_db(conn)
    conn.execute(
        "INSERT INTO accounts(account_number) VALUES ('caller-owned')"
    )

    def fail(*args):
        raise RuntimeError('forced failure')

    monkeypatch.setattr(classify, 'apply_rules', fail)
    try:
        with pytest.raises(RuntimeError, match='forced failure'):
            import_statement(conn, path, path.name)
        assert conn.execute(
            'SELECT account_number FROM accounts'
        ).fetchone()[0] == 'caller-owned'
        assert conn.execute('SELECT COUNT(*) FROM imports').fetchone()[0] == 0
        conn.commit()
        assert conn.execute(
            'SELECT COUNT(*) FROM accounts'
        ).fetchone()[0] == 1
    finally:
        conn.close()


def test_preview_is_read_only_and_requires_warning_approval(tmp_path) -> None:
    path = tmp_path / 'preview-warning.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    wb.active['A3'] = 'not-a-date'
    wb.save(path)
    wb.close()
    conn = get_conn(':memory:')
    init_db(conn)
    try:
        preview = preview_statement(conn, path, path.name)
        assert preview.masked_account_number.endswith('0001')
        assert 'MN030005005000000001' not in preview.masked_account_number
        assert preview.rows_new == 11
        assert preview.rows_duplicate == 0
        assert preview.rows_skipped == 1
        assert preview.income_total_minor > 0
        assert preview.expense_total_minor > 0
        assert len(preview.sample_rows) == 5
        assert preview.requires_confirmation
        for table in ('accounts', 'imports', 'transactions'):
            assert conn.execute(
                'SELECT COUNT(*) FROM ' + table
            ).fetchone()[0] == 0
        with pytest.raises(ValueError, match='explicit confirmation'):
            import_preview(conn, preview)
        assert conn.execute('SELECT COUNT(*) FROM imports').fetchone()[0] == 0
        result = import_preview(conn, preview, accept_warnings=True)
        assert result.rows_new == 11
        assert result.rows_skipped == 1
        assert conn.execute('SELECT COUNT(*) FROM transactions').fetchone()[0] == 11
    finally:
        conn.close()

def test_preview_duplicate_counts_and_manual_data_survive_reimport(tmp_path) -> None:
    path = tmp_path / 'duplicate-preview.xlsx'
    make_fixture(path)
    conn = get_conn(':memory:')
    init_db(conn)
    try:
        first = import_statement(conn, path, path.name)
        txn = conn.execute(
            "SELECT id, direction FROM transactions LIMIT 1"
        ).fetchone()
        category_name = 'Groceries / Food' if txn['direction'] == 'out' else 'Allowance / Family support'
        category_id = conn.execute(
            'SELECT id FROM categories WHERE name=?', (category_name,)
        ).fetchone()[0]
        set_category(conn, [txn['id']], category_id)
        conn.execute(
            'UPDATE transactions SET note=? WHERE id=?', ('keep this note', txn['id'])
        )
        conn.commit()
        preview = preview_statement(conn, path, path.name)
        assert preview.rows_new == 0
        assert preview.rows_duplicate == 12
        assert preview.requires_confirmation
        second = import_preview(conn, preview, accept_warnings=True)
        assert (second.rows_new, second.rows_duplicate) == (0, 12)
        preserved = conn.execute(
            'SELECT category_id, classified_by, note FROM transactions WHERE id=?',
            (txn['id'],),
        ).fetchone()
        assert tuple(preserved) == (category_id, 'manual', 'keep this note')
        assert conn.execute('SELECT COUNT(*) FROM transactions').fetchone()[0] == 12
    finally:
        conn.close()

def test_fatal_statement_is_not_saved(tmp_path) -> None:
    path = tmp_path / 'unrecognized.xlsx'
    from openpyxl import Workbook

    wb = Workbook()
    wb.active['A1'] = 'not a bank statement'
    wb.save(path)
    wb.close()
    conn = get_conn(':memory:')
    init_db(conn)
    try:
        with pytest.raises(ValueError, match='Unrecognized'):
            preview_statement(conn, path, path.name)
        assert conn.execute('SELECT COUNT(*) FROM imports').fetchone()[0] == 0
        assert conn.execute('SELECT COUNT(*) FROM transactions').fetchone()[0] == 0
    finally:
        conn.close()

def test_statement_preview_closes_workbook(tmp_path, monkeypatch) -> None:
    from core import ingest

    path = tmp_path / 'closed.xlsx'
    make_fixture(path)
    actual_loader = ingest.load_workbook
    observed = {'closed': False}

    def tracking_loader(*args, **kwargs):
        workbook = actual_loader(*args, **kwargs)
        original_close = workbook.close

        def close():
            observed['closed'] = True
            original_close()

        workbook.close = close
        return workbook

    monkeypatch.setattr(ingest, 'load_workbook', tracking_loader)
    conn = get_conn(':memory:')
    init_db(conn)
    try:
        preview_statement(conn, path, path.name)
        assert observed['closed']
    finally:
        conn.close()

def test_import_preview_failure_rolls_back_all_writes(tmp_path, monkeypatch) -> None:
    from core import classify

    path = tmp_path / 'preview-rollback.xlsx'
    make_fixture(path)
    conn = get_conn(':memory:')
    init_db(conn)

    def fail(*args):
        raise RuntimeError('forced preview import failure')

    monkeypatch.setattr(classify, 'apply_rules', fail)
    try:
        preview = preview_statement(conn, path, path.name)
        with pytest.raises(RuntimeError, match='forced preview import failure'):
            import_preview(conn, preview)
        for table in ('accounts', 'imports', 'transactions'):
            assert conn.execute(
                'SELECT COUNT(*) FROM ' + table
            ).fetchone()[0] == 0
    finally:
        conn.close()

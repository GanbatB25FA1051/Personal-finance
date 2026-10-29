"""Check detection, validation and merchant normalization."""
from pathlib import Path

import pytest
from openpyxl import load_workbook
from banks.base import detect_parser
from banks.khan_bank import KhanBankParser
from tests.make_fixture import make_fixture


def test_fixture(tmp_path: Path) -> None:
    path = tmp_path / 'statement.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    parser = detect_parser(wb)
    assert isinstance(parser, KhanBankParser)
    result = parser.parse(wb)
    assert len(result.rows) == 12
    assert result.warnings == []
    assert result.account_number == 'MN030005005000000001'
    assert result.rows[0].amount == -1520000
    assert result.rows[5].amount == 1758528
    assert sum(max(r.amount, 0) for r in result.rows) == result.control_credit
    assert sum(min(r.amount, 0) for r in result.rows) == result.control_debit
    wb.active['F3'] = 0
    assert parser.parse(wb).warnings
    wb.close()


@pytest.mark.parametrize(('raw', 'expected'), [
    ('qpay fino-940482663759306, Дэвжих супермаркет', 'дэвжих супермаркет'),
    ('QPAY 243969128072217, 2851837,DODO PI', 'dodo pi'),
    ('qpay 414170975548256 UB PROSPERITY FOOD LLC 28500', 'ub prosperity food llc'),
    ('Z,30022607125498,2509015047,261002225943899', 'z-recurring'),
    ('qpay 123456789, Qpay, Qmenu, Gochu - Их Дэлг', 'gochu - их дэлг'),
])
def test_clean_merchant(raw: str, expected: str) -> None:
    assert KhanBankParser().clean_merchant(raw) == expected


@pytest.mark.parametrize(('raw', 'expected'), [
    ('qpay fino fee-424872322362495, шимтгэл', True),
    ('Ухаалаг банкны гүйлгээний хураамж', True),
    ('qpay 123456789 шимтгэл', True), ('DODO PI', False), ('A', False),
])
def test_is_fee(raw: str, expected: bool) -> None:
    assert KhanBankParser().is_fee(raw) is expected


def test_shifted_and_reordered_headers(tmp_path: Path) -> None:
    path = tmp_path / 'shifted.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    ws = wb.active
    ws.insert_rows(2)
    for row in ws.iter_rows(min_row=3):
        row[0].value, row[6].value = row[6].value, row[0].value
    parser = detect_parser(wb)
    result = parser.parse(wb)
    assert len(result.rows) == 12
    assert result.warnings == []
    assert result.rows[0].amount == -1520000
    wb.close()


def test_footer_totals_follow_reordered_columns_and_header_can_be_shifted(
    tmp_path: Path,
) -> None:
    path = tmp_path / 'reordered-footer.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    ws = wb.active
    ws.insert_rows(2, 4)
    order = [7, 0, 3, 4, 5, 2, 6, 1]
    for row_number in range(6, ws.max_row + 1):
        old = [ws.cell(row_number, col).value for col in range(1, 9)]
        for new_col, old_col in enumerate(order, 1):
            ws.cell(row_number, new_col, old[old_col])
    parser = detect_parser(wb)
    assert isinstance(parser, KhanBankParser)
    result = parser.parse(wb)
    assert result.fatal_errors == []
    assert len(result.rows) == 12
    assert result.control_credit == sum(max(row.amount, 0) for row in result.rows)
    assert result.control_debit == sum(min(row.amount, 0) for row in result.rows)
    assert result.warnings == []
    wb.close()

def test_positive_debit_is_normalized_and_requires_review(tmp_path: Path) -> None:
    path = tmp_path / 'positive-debit.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    wb.active['E3'] = 15200
    result = KhanBankParser().parse(wb)
    assert result.rows[0].amount == -1_520_000
    assert any('positive debit normalized' in warning for warning in result.warnings)
    wb.close()

def test_invalid_transaction_row_is_counted_as_skipped(tmp_path: Path) -> None:
    path = tmp_path / 'invalid-row.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    wb.active['A3'] = 'not-a-date'
    result = KhanBankParser().parse(wb)
    assert len(result.rows) == 11
    assert result.rows_skipped == 1
    assert any('row skipped' in warning for warning in result.warnings)
    wb.close()


def test_control_total_mismatch_is_a_review_warning(tmp_path: Path) -> None:
    path = tmp_path / 'control-mismatch.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    ws = wb.active
    footer = next(
        row for row in range(1, ws.max_row + 1)
        if str(ws.cell(row, 1).value or '').startswith('Нийт дүн:')
    )
    ws.cell(footer, 4, '0.00')
    result = KhanBankParser().parse(wb)
    assert any('Total credit differs' in warning for warning in result.warnings)
    wb.close()

def test_balance_mismatch_is_a_review_warning(tmp_path: Path) -> None:
    path = tmp_path / 'balance-mismatch.xlsx'
    make_fixture(path)
    wb = load_workbook(path)
    wb.active['F3'] = 0
    result = KhanBankParser().parse(wb)
    assert any('opening balance plus amount' in warning for warning in result.warnings)
    wb.close()
